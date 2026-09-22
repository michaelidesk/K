import sys, datetime
sys.path.insert(0, "/tmp/claude-0/-home-user-K/c282222a-9a6f-5b9c-b93f-bdc4098c8842/scratchpad")
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.comments import Comment
from internal_data import RAISE_PARAMS, EXISTING_PIPELINE
from external_data import (MARKET_METRICS, SEGMENTS, UNIVERSE, CONTACTS, CASCADE_MAP,
                           CONFERENCES, SOURCES, OUTREACH_WAVES, MESSAGE_ANGLES)

OUT = "/home/user/K/deliverables/Impactus_EU_Private_Credit_Investor_Map_Sep2026.xlsx"

NAVY = "1F2A44"; GREY = "F2F2F2"; MID = "D9DEE8"; WHITE = "FFFFFF"; YELLOW = "FFF2CC"
BLUE_INPUT = Font(name="Arial", size=10, color="0000FF")
BLACK = Font(name="Arial", size=10, color="000000")
GREEN_LINK = Font(name="Arial", size=10, color="008000")
BOLD = Font(name="Arial", size=10, bold=True)
HDR = Font(name="Arial", size=10, bold=True, color=WHITE)
TITLE = Font(name="Arial", size=14, bold=True, color=NAVY)
SUB = Font(name="Arial", size=10, italic=True, color="595959")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
EURM = '"€"#,##0.0"m";("€"#,##0.0"m");"-"'
EUR = '"€"#,##0;("€"#,##0);"-"'
PCT = '0.0%'
NUM = '#,##0;(#,##0);"-"'
X = '0.00"x"'

DISCLAIMER = ("AIFMD Art. 23 disclosure: Impactus RAIF V.C.I.C. Plc is a Registered Alternative Investment Fund (CySEC RAIF16) "
              "externally managed by Argus Management Ltd (CySEC AIFM 44/56/2013). This workbook is an internal working document "
              "prepared for the Sponsor and is not an offer, solicitation or marketing communication. Interests in the Fund and its "
              "sub-funds (Impactus Private Equity Fund, CPEF I Curity, Impactus Pyrolysis / IVP) are available only to professional "
              "and well-informed investors within the meaning of AIFMD Art. 4(1)(ag), MiFID II and the Cyprus AIF Law 124(I)/2018. "
              "Prospective investors must read the AIFMD Art. 23 disclosure document, the relevant offering supplement and term sheet "
              "before investing. Targets are gross, forward-looking and not guaranteed; capital is at risk. Third-party names are used "
              "for identification only and do not imply endorsement.")

def hdr_row(ws, row, headers, widths=None, fill=NAVY):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = HDR; c.fill = PatternFill("solid", fgColor=fill); c.alignment = CENTER; c.border = BORDER
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i-1]
    ws.row_dimensions[row].height = 30

def title(ws, text, subtitle=None):
    ws["A1"] = text; ws["A1"].font = TITLE
    if subtitle:
        ws["A2"] = subtitle; ws["A2"].font = SUB
    ws.sheet_view.showGridLines = False

def write_rows(ws, start_row, rows, fmts=None, wrap_cols=None):
    r = start_row
    for row in rows:
        for j, v in enumerate(row, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = BLACK; c.border = BORDER; c.alignment = WRAP
            if fmts and j in fmts and isinstance(v, (int, float)):
                c.number_format = fmts[j]
        r += 1
    return r

def inp(ws, ref, value, fmt=None, note=None):
    c = ws[ref]; c.value = value; c.font = BLUE_INPUT; c.fill = PatternFill("solid", fgColor=YELLOW); c.border = BORDER
    if fmt: c.number_format = fmt
    if note: c.comment = Comment(note, "Impactus analysis")
    return c

def fml(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]; c.value = formula; c.font = BOLD if bold else BLACK; c.border = BORDER
    if fmt: c.number_format = fmt
    return c

def label(ws, ref, text, bold=False):
    c = ws[ref]; c.value = text; c.font = BOLD if bold else BLACK; c.alignment = WRAP; c.border = BORDER
    return c

# ---------------------------------------------------------------- Cover
def sheet_cover(wb):
    ws = wb.active; ws.title = "Cover"
    title(ws, "Europe Private Debt & Alternative Credit Investors - Relevance to the CPEF I Curity and Impactus Pyrolysis (IPL) raises",
          f"Impactus Private Equity Group - internal working workbook - prepared {datetime.date.today().strftime('%d %B %Y')} - source article: Cascade Debt, 'Private Debt and Alternative Credit Investors in Europe' (2026 Europe Private ABF Investor Market Map)")
    ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 120
    rows = [
        ("Purpose", "Map the European private debt / alternative credit investor universe (Cascade 2026 ABF map plus the wider direct lending, special situations, infrastructure credit, venture debt, impact and DFI universe), score each firm for relevance to the IPL (€5m equity + €35m venture debt + senior refinancing optionality) and CPEF (8% note today; €3m to €10m growth facility next) raises, and give a named, channel-specific contact plan."),
        ("How to read the tabs", "1 Exec_Summary: headline counts and capacity. 2 Market_Overview: sourced market data. 3 Segment_Map: which credit segment fits which Impactus need. 4 Cascade_Map_2026: the article's investor map transcribed and scored. 5 Investor_Universe: scored target list (firm level). 6 Contacts: named decision-makers with channel recommendation. 7 Existing_Pipeline: warm relationships already in Impactus records. 8 IPL_Debt_Capacity and 9 CPEF_Debt_Capacity: live models. 10 Outreach_Plan and 11 Conferences. 12 Scoring_Weights: change weights here. 13 Sources."),
        ("Colour legend", "Blue text on yellow = input you can change. Black = formula or fixed data. Green = link to another tab. Scores recalculate when weights on Scoring_Weights change."),
        ("Data status", "External figures are sourced (see Sources tab, URL per row). Contact emails: only those published on public pages are shown as 'published'; 'pattern - unverified' means the firm's public address format, to be verified before sending. LinkedIn URLs are public profile URLs found in September 2026. Apollo.io enrichment was not available in this session (free plan blocks API search and enrichment), so no purchased data is included."),
        ("Analytical stance", "Bear case first. Rows flagged 'DD challenge' name the objection an LP due-diligence lead, lender credit committee or CySEC reviewer would raise. The memo (Word) carries the argued position; this workbook carries the evidence and the working models."),
        ("Disclaimer", DISCLAIMER),
    ]
    r = 4
    for k, v in rows:
        ws.cell(row=r, column=1, value=k).font = BOLD
        c = ws.cell(row=r, column=2, value=v); c.font = BLACK; c.alignment = WRAP
        ws.row_dimensions[r].height = 15 * max(2, int(len(v) / 110) + 1)
        r += 1
    return ws

# ---------------------------------------------------------------- Raise params + pipeline
def sheet_raise(wb):
    ws = wb.create_sheet("Raise_Parameters")
    title(ws, "Raise parameters and internal reference points (Sponsor sources, September 2026)", "All figures as they stand in the Sponsor's own working documents; PMO 02-09-2026 reset flagged where relevant")
    hdr_row(ws, 4, ["Item", "Value", "Unit", "Source", "Comment / DD flag"], [46, 40, 8, 38, 70])
    write_rows(ws, 5, RAISE_PARAMS, fmts={2: '#,##0.0'})
    ws.freeze_panes = "A5"
    return ws

def sheet_pipeline(wb):
    ws = wb.create_sheet("Existing_Pipeline")
    title(ws, "Existing and warm relationships already in Impactus records (Outlook, SharePoint tracker, Fathom calls)", "Use these paths before any cold outreach - a warm introduction converts 5x to 10x better than a cold LinkedIn message (see Outreach_Plan)")
    hdr_row(ws, 4, ["Counterparty", "Type", "Owner", "Status / last step", "Relevance", "Suggested next action", "Source"], [46, 26, 12, 55, 18, 55, 34])
    write_rows(ws, 5, EXISTING_PIPELINE)
    ws.freeze_panes = "B5"
    return ws

# ---------------------------------------------------------------- IPL debt capacity model
def sheet_ipl_model(wb):
    ws = wb.create_sheet("IPL_Debt_Capacity")
    title(ws, "IPL / Thermo Lysi - what the credit market can carry, and what is left for equity", "Live model: change blue inputs. Base inputs per PwC Greece model Iteration 3 v.6 (2031 full year) and Credia LOI; PMO reset items flagged.")
    for col, w in zip("ABCDEFGH", [44, 16, 16, 16, 16, 16, 16, 40]):
        ws.column_dimensions[col].width = w
    # Inputs
    label(ws, "A4", "OPERATING INPUTS (2031 full year)", True)
    label(ws, "A5", "Throughput (t)"); inp(ws, "B5", 44776, NUM, "PwC model Iteration 3 v.6 - 2031 throughput")
    label(ws, "A6", "rCB yield"); inp(ws, "B6", 0.303, PCT)
    label(ws, "A7", "rCB price (€/t)"); inp(ws, "B7", 1500, EUR, "Pyrum letter / Continental framework; Sergo call Jun 2025 cited a €1,350/t floor")
    label(ws, "A8", "Pyrolysis oil yield"); inp(ws, "B8", 0.20, PCT)
    label(ws, "A9", "Oil price (€/t)"); inp(ws, "B9", 845, EUR, "PwC model; BASF HoT. Jun 2025 call cited €515/t floor - stress case")
    label(ws, "A10", "Steel yield"); inp(ws, "B10", 0.245, PCT)
    label(ws, "A11", "Steel price (€/t)"); inp(ws, "B11", 250, EUR)
    label(ws, "A12", "Gate fee (€/t) on all throughput"); inp(ws, "B12", 100, EUR, "MODEL assumption. Contracted today: €62/t on 11ktpa only (PMO MoM 03-09-2026). This is the swing item.")
    label(ws, "A13", "Carbon credits (€/t)"); inp(ws, "B13", 25, EUR)
    label(ws, "A14", "Cash opex (€/t)"); inp(ws, "B14", 317, EUR, "Derived: (€36.1m revenue - €21.9m EBITDA) / 44,776t per PwC model")
    label(ws, "A15", "Maintenance capex (€/t)"); inp(ws, "B15", 20, EUR, "Assumption - not in Snapshot; needs verification against PwC model")
    label(ws, "A16", "Corporate tax rate (Greece)"); inp(ws, "B16", 0.22, PCT, "Greek CIT 22%")
    label(ws, "A18", "SENIOR DEBT INPUTS (Credia LOI, Aug 2026)", True)
    label(ws, "A19", "Interest rate (all-in)"); inp(ws, "B19", 0.065, PCT, "Credia LOI: 6.5% p.a. indicative")
    label(ws, "A20", "Repayment tenor (years, post availability)"); inp(ws, "B20", 10, NUM, "4 + 10 years")
    label(ws, "A21", "Minimum DSCR (sizing)"); inp(ws, "B21", 1.5, X, "Credia sizing 1.5x; LOI average DSCR 1.3x")
    label(ws, "A22", "Senior facility in Sponsor plan (€m)"); inp(ws, "B22", 77.4, EURM)
    label(ws, "A24", "VENTURE DEBT / PRIVATE CREDIT TRANCHE", True)
    label(ws, "A25", "Tranche size (€m)"); inp(ws, "B25", 35, EURM)
    label(ws, "A26", "Cash coupon"); inp(ws, "B26", 0.04, PCT)
    label(ws, "A27", "PIK coupon"); inp(ws, "B27", 0.05, PCT)
    label(ws, "A28", "Warrant value (% of tranche, PV)"); inp(ws, "B28", 0.03, PCT, "Assumption - 3% of tranche; typical VD warrant coverage 5-15% of loan translates to 1-4% PV")
    label(ws, "A29", "Term (years)"); inp(ws, "B29", 7, NUM)
    label(ws, "A31", "OTHER SOURCES (€m)", True)
    label(ws, "A32", "EU Innovation Fund grant"); inp(ws, "B32", 29.4, EURM)
    label(ws, "A33", "Shareholder equity committed"); inp(ws, "B33", 15, EURM)
    label(ws, "A34", "Shareholder loans (10% simple)"); inp(ws, "B34", 6, EURM)
    label(ws, "A35", "VAT facility"); inp(ws, "B35", 13.3, EURM)
    label(ws, "A36", "Total sources required (Sponsor plan)"); inp(ws, "B36", 176, EURM, "Snapshot p.9: plant capex €118.5m + development, fees, WC, reserves")
    label(ws, "A37", "Capex stress (PMO high case €150m plant vs €118.5m) - uplift"); inp(ws, "B37", 31.5, EURM, "PMO MoM 03-09-2026: '€120m or €150m' unresolved")
    label(ws, "A38", "Target equity IRR"); inp(ws, "B38", 0.22, PCT)

    # Outputs
    label(ws, "D4", "OUTPUTS", True)
    label(ws, "D5", "rCB revenue (€m)"); fml(ws, "E5", "=B5*B6*B7/1000000", EURM)
    label(ws, "D6", "Oil revenue (€m)"); fml(ws, "E6", "=B5*B8*B9/1000000", EURM)
    label(ws, "D7", "Steel revenue (€m)"); fml(ws, "E7", "=B5*B10*B11/1000000", EURM)
    label(ws, "D8", "Gate fee revenue (€m)"); fml(ws, "E8", "=B5*B12/1000000", EURM)
    label(ws, "D9", "Carbon credit revenue (€m)"); fml(ws, "E9", "=B5*B13/1000000", EURM)
    label(ws, "D10", "Total revenue (€m)"); fml(ws, "E10", "=SUM(E5:E9)", EURM, True)
    label(ws, "D11", "Cash opex (€m)"); fml(ws, "E11", "=B5*B14/1000000", EURM)
    label(ws, "D12", "EBITDA (€m)"); fml(ws, "E12", "=E10-E11", EURM, True)
    label(ws, "D13", "EBITDA margin"); fml(ws, "E13", "=IF(E10=0,0,E12/E10)", PCT)
    label(ws, "D14", "CFADS proxy (€m) = EBITDA x (1 - tax) - maintenance capex"); fml(ws, "E14", "=E12*(1-B16)-B5*B15/1000000", EURM, True)
    ws["D14"].comment = Comment("Simplification: ignores tax shield from depreciation and interest; conservative for sizing.", "Impactus analysis")
    label(ws, "D16", "Senior debt capacity at DSCR (€m)"); fml(ws, "E16", "=PV(B19,B20,-E14/B21)", EURM, True)
    ws["D16"].comment = Comment("Annuity-based: PV of CFADS/DSCR over the repayment tenor at the facility rate.", "Impactus analysis")
    label(ws, "D17", "Senior capacity vs Sponsor plan (€m)"); fml(ws, "E17", "=E16-B22", EURM)
    label(ws, "D18", "Senior capacity at 1.3x DSCR (LOI average) (€m)"); fml(ws, "E18", "=PV(B19,B20,-E14/1.3)", EURM)
    label(ws, "D20", "VD all-in cost (cash + PIK + warrant amortised)"); fml(ws, "E20", "=B26+B27+B28/B29", PCT, True)
    label(ws, "D21", "VD annual cash interest (€m)"); fml(ws, "E21", "=B25*B26", EURM)
    label(ws, "D22", "VD balance at maturity with PIK (€m)"); fml(ws, "E22", "=B25*(1+B27)^B29", EURM)
    label(ws, "D24", "Sources check (€m): grant + equity + SHL + VAT + senior plan + VD"); fml(ws, "E24", "=B32+B33+B34+B35+B22+B25", EURM, True)
    label(ws, "D25", "Gap vs total sources required (€m)"); fml(ws, "E25", "=E24-B36", EURM, True)
    label(ws, "D26", "Gap if capex stress applies (€m)"); fml(ws, "E26", "=E25-B37", EURM, True)
    label(ws, "D27", "Gap if senior sized on model CFADS instead of plan (€m)"); fml(ws, "E27", "=E25+E17", EURM, True)
    label(ws, "D29", "Blended cost of capital (senior, VD, SHL, equity)"); fml(ws, "E29", "=(B22*B19+B25*E20+B34*0.10+B33*B38)/(B22+B25+B34+B33)", PCT, True)
    label(ws, "D30", "Interest cover: EBITDA / (senior cash interest + VD cash interest)"); fml(ws, "E30", "=IF((B22*B19+E21)=0,0,E12/(B22*B19+E21))", X, True)

    # Sensitivity: gate fee x rCB price -> senior capacity
    label(ws, "A41", "SENSITIVITY - senior debt capacity (€m) by gate fee (rows, €/t) and rCB price (columns, €/t), all other inputs as above", True)
    prices = [1000, 1250, 1500, 1750]
    fees = [0, 30, 62, 100, 130]
    ws["A42"].value = "Gate fee €/t \\ rCB €/t"; ws["A42"].font = BOLD; ws["A42"].border = BORDER
    for j, p in enumerate(prices):
        c = ws.cell(row=42, column=2 + j, value=p); c.font = BOLD; c.number_format = EUR; c.border = BORDER; c.fill = PatternFill("solid", fgColor=MID)
    for i, f in enumerate(fees):
        r = 43 + i
        c = ws.cell(row=r, column=1, value=f); c.font = BOLD; c.number_format = EUR; c.border = BORDER; c.fill = PatternFill("solid", fgColor=MID)
        for j, p in enumerate(prices):
            col = get_column_letter(2 + j)
            ebitda = f"(($B$5*$B$6*{col}$42+$B$5*$B$8*$B$9+$B$5*$B$10*$B$11+$B$5*$A{r}+$B$5*$B$13)/1000000-$E$11)"
            cfads = f"({ebitda}*(1-$B$16)-$B$5*$B$15/1000000)"
            cell = ws.cell(row=r, column=2 + j, value=f"=PV($B$19,$B$20,-{cfads}/$B$21)")
            cell.number_format = EURM; cell.border = BORDER; cell.font = BLACK
    ws.conditional_formatting.add("B43:E47", ColorScaleRule(start_type="min", start_color="F8CBAD", mid_type="percentile", mid_value=50, mid_color="FFEB9C", end_type="max", end_color="C6EFCE"))
    label(ws, "A49", "Reading: the contracted position today (€62/t on 11ktpa, i.e. c.€15/t on 45ktpa) sits between the 0 and 30 rows. The Sponsor plan senior of €77.4m requires the €100/t row at €1,500/t rCB or better. Every €10/t of gate fee on 45ktpa is c.€0.45m EBITDA and c.€1.7m of senior capacity at 6.5% / 10y / 1.5x.")
    ws["A49"].alignment = WRAP; ws.merge_cells("A49:H50"); ws.row_dimensions[49].height = 30
    # Sensitivity 2: EBITDA by oil price x rCB price
    label(ws, "A52", "SENSITIVITY - EBITDA (€m) by oil price (rows, €/t) and rCB price (columns, €/t)", True)
    oils = [515, 650, 845, 1000]
    ws["A53"].value = "Oil €/t \\ rCB €/t"; ws["A53"].font = BOLD; ws["A53"].border = BORDER
    for j, p in enumerate(prices):
        c = ws.cell(row=53, column=2 + j, value=p); c.font = BOLD; c.number_format = EUR; c.border = BORDER; c.fill = PatternFill("solid", fgColor=MID)
    for i, o in enumerate(oils):
        r = 54 + i
        c = ws.cell(row=r, column=1, value=o); c.font = BOLD; c.number_format = EUR; c.border = BORDER; c.fill = PatternFill("solid", fgColor=MID)
        for j, p in enumerate(prices):
            col = get_column_letter(2 + j)
            cell = ws.cell(row=r, column=2 + j, value=f"=($B$5*$B$6*{col}$53+$B$5*$B$8*$A{r}+$B$5*$B$10*$B$11+$B$5*$B$12+$B$5*$B$13)/1000000-$E$11")
            cell.number_format = EURM; cell.border = BORDER; cell.font = BLACK
    ws.conditional_formatting.add("B54:E57", ColorScaleRule(start_type="min", start_color="F8CBAD", mid_type="percentile", mid_value=50, mid_color="FFEB9C", end_type="max", end_color="C6EFCE"))
    label(ws, "A59", "DD challenge (lender credit committee): the model EBITDA of €21.9m rests on €1,500/t rCB and €100/t gate fee across 45ktpa; neither is contracted at that level today. A credit committee will size on contracted cash flows only, which is why the venture-debt / private-credit tranche exists - it is the price of the gap between contracted and modelled revenue, not a cheaper alternative to equity.")
    ws["A59"].alignment = WRAP; ws.merge_cells("A59:H61"); ws.row_dimensions[59].height = 30
    return ws

# ---------------------------------------------------------------- CPEF model
def sheet_cpef_model(wb):
    ws = wb.create_sheet("CPEF_Debt_Capacity")
    title(ws, "CPEF I Curity / Curity Pharma - how much growth debt the asset can carry, and at what price", "Live model. Base inputs from the Meerkat valuation business plan (stressed case) and the June 2026 Investor Pack; refresh with FY2025 actuals before external use.")
    for col, w in zip("ABCDEFG", [46, 16, 16, 16, 16, 16, 40]):
        ws.column_dimensions[col].width = w
    label(ws, "A4", "OPERATING INPUTS", True)
    label(ws, "A5", "Grams sold p.a."); inp(ws, "B5", 3498219, NUM, "Meerkat stressed plan: 3.5m g at capacity (Phase 1)")
    label(ws, "A6", "Net price (€/g)"); inp(ws, "B6", 3.00, '"€"0.00', "Stressed €3.00/g; base case €4.85/g (Sponsor). German wholesale c.€5/g in 2025 per Business of Cannabis; verify current EU-GMP flower prices")
    label(ws, "A7", "EBITDA margin"); inp(ws, "B7", 0.50, PCT, "Meerkat plan: 50-51%")
    label(ws, "A8", "Tax rate (Greece)"); inp(ws, "B8", 0.22, PCT)
    label(ws, "A9", "Maintenance capex (€m)"); inp(ws, "B9", 0.11, EURM, "Meerkat plan c.€0.1m p.a.")
    label(ws, "A10", "Existing senior bank debt (€m)"); inp(ws, "B10", 5.57, EURM, "Greek bank facility per plan; confirm drawn balance")
    label(ws, "A11", "Existing bank rate"); inp(ws, "B11", 0.06, PCT, "Assumption")
    label(ws, "A13", "GROWTH FACILITY INPUTS", True)
    label(ws, "A14", "Target leverage (net debt / EBITDA)"); inp(ws, "B14", 2.5, X, "Specialist cannabis lenders: 2-3x; mainstream EU lower-mid-market direct lenders: would not lend (sector exclusion) - see Investor_Universe")
    label(ws, "A15", "Growth facility rate (cash)"); inp(ws, "B15", 0.13, PCT, "Cannabis credit pricing 12-16% (Chicago Atlantic / AFC disclosed portfolio yields); EU venture debt 10-13% + warrants")
    label(ws, "A16", "Tenor (years)"); inp(ws, "B16", 4, NUM)
    label(ws, "A17", "Warrants / fees (% of facility, PV)"); inp(ws, "B17", 0.04, PCT)
    label(ws, "A18", "CPEF note coupon (reference)"); inp(ws, "B18", 0.08, PCT)
    label(ws, "A19", "Curity Pharma last-marked equity value (€m)"); inp(ws, "B19", 22.2, EURM)
    label(ws, "A20", "CPEF holding in Curity"); inp(ws, "B20", 0.3375, PCT)

    label(ws, "D4", "OUTPUTS", True)
    label(ws, "D5", "Revenue (€m)"); fml(ws, "E5", "=B5*B6/1000000", EURM, True)
    label(ws, "D6", "EBITDA (€m)"); fml(ws, "E6", "=E5*B7", EURM, True)
    label(ws, "D7", "Total debt capacity at target leverage (€m)"); fml(ws, "E7", "=E6*B14", EURM, True)
    label(ws, "D8", "Headroom for growth facility after existing debt (€m)"); fml(ws, "E8", "=MAX(0,E7-B10)", EURM, True)
    label(ws, "D9", "Annual cash interest on growth facility (€m)"); fml(ws, "E9", "=E8*B15", EURM)
    label(ws, "D10", "Total cash interest (existing + growth) (€m)"); fml(ws, "E10", "=B10*B11+E9", EURM)
    label(ws, "D11", "Interest cover (EBITDA / cash interest)"); fml(ws, "E11", "=IF(E10=0,0,E6/E10)", X, True)
    label(ws, "D12", "Free cash after tax, capex and interest (€m)"); fml(ws, "E12", "=E6*(1-B8)-B9-E10", EURM, True)
    label(ws, "D13", "Debt service capacity (FCF / (growth facility / tenor))"); fml(ws, "E13", "=IF(E8=0,0,E12/(E8/B16))", X)
    label(ws, "D14", "All-in cost of growth facility"); fml(ws, "E14", "=B15+B17/B16", PCT, True)
    label(ws, "D15", "Spread of growth facility over CPEF 8% note"); fml(ws, "E15", "=E14-B18", PCT)
    label(ws, "D17", "Look-through: CPEF share of Curity equity value (€m)"); fml(ws, "E17", "=B19*B20", EURM)
    label(ws, "D18", "Implied EV / EBITDA at last mark (with existing debt)"); fml(ws, "E18", "=IF(E6=0,0,(B19+B10)/E6)", X)

    label(ws, "A23", "SENSITIVITY - growth-facility headroom (€m) by price €/g (rows) and grams sold (columns)", True)
    grams = [1500000, 2500000, 3500000, 5100000]
    px = [2.50, 3.00, 3.75, 4.85]
    ws["A24"].value = "€/g \\ grams"; ws["A24"].font = BOLD; ws["A24"].border = BORDER
    for j, g in enumerate(grams):
        c = ws.cell(row=24, column=2 + j, value=g); c.font = BOLD; c.number_format = NUM; c.border = BORDER; c.fill = PatternFill("solid", fgColor=MID)
    for i, p in enumerate(px):
        r = 25 + i
        c = ws.cell(row=r, column=1, value=p); c.font = BOLD; c.number_format = '"€"0.00'; c.border = BORDER; c.fill = PatternFill("solid", fgColor=MID)
        for j, g in enumerate(grams):
            col = get_column_letter(2 + j)
            cell = ws.cell(row=r, column=2 + j, value=f"=MAX(0,{col}$24*$A{r}/1000000*$B$7*$B$14-$B$10)")
            cell.number_format = EURM; cell.border = BORDER; cell.font = BLACK
    ws.conditional_formatting.add("B25:E28", ColorScaleRule(start_type="min", start_color="F8CBAD", mid_type="percentile", mid_value=50, mid_color="FFEB9C", end_type="max", end_color="C6EFCE"))
    label(ws, "A30", "DD challenge (LP due-diligence lead): Curity is early-revenue (first exports 2026, Galenica distribution from October 2026). Until twelve months of audited sales exist, the only lenders are specialist cannabis credit funds at 12-16% or an asset-based lender against EU-GMP inventory and receivables. The 8% note is therefore priced below where a third-party lender would price the same risk; the difference is being carried by the Sponsor's relationship with its own investor base.")
    ws["A30"].alignment = WRAP; ws.merge_cells("A30:G32"); ws.row_dimensions[30].height = 30
    return ws

# ---------------------------------------------------------------- Scoring weights
def sheet_weights(wb):
    ws = wb.create_sheet("Scoring_Weights")
    title(ws, "Scoring weights - change here; Investor_Universe and Cascade_Map_2026 recalculate", "Each criterion is scored 0 to 5 per firm; weighted score = SUMPRODUCT(scores, weights) / SUM(weights) x 20 -> 0 to 100")
    ws.column_dimensions["A"].width = 40; ws.column_dimensions["B"].width = 12; ws.column_dimensions["C"].width = 90
    hdr_row(ws, 4, ["Criterion", "Weight", "What 5 means"])
    crit = [
        ("Strategy fit - IPL (hard-asset / infra / FOAK credit)", 25, "Runs a strategy that lends to or invests in single-asset industrial / energy-transition projects at construction or pre-completion stage"),
        ("Strategy fit - CPEF (pharma / cannabis / growth credit)", 15, "Lends to regulated pharma-adjacent or cannabis producers, or runs healthcare growth credit with no cannabis exclusion"),
        ("Ticket fit (€5m to €40m)", 15, "Minimum ticket at or below €10m and sweet spot inside €5m to €40m"),
        ("Geography (Greece / Cyprus / Southern Europe)", 15, "Has an office, portfolio or repeated deal history in Greece or Cyprus"),
        ("Risk appetite (FOAK, PIK, warrants, sub-IG)", 10, "Documented appetite for first-of-a-kind technology, PIK, warrants or construction risk"),
        ("Warm path (existing Impactus relationship or one-step intro)", 10, "Already in the Impactus pipeline, or reachable via Argus, Venero, Sirec, KPMG, PwC, Axia, Pyrum, Archirodon, EIB"),
        ("Speed / decision process", 10, "Documented ability to close in under six months; small IC; local decision-maker"),
    ]
    r = 5
    for name, w, desc in crit:
        label(ws, f"A{r}", name); inp(ws, f"B{r}", w, NUM); label(ws, f"C{r}", desc)
        r += 1
    label(ws, "A12", "Total weight", True); fml(ws, "B12", "=SUM(B5:B11)", NUM, True)
    label(ws, "E4", "Helper row (weights transposed for SUMPRODUCT - do not edit)")
    for j in range(7):
        fml(ws, f"{get_column_letter(5 + j)}5", f"=B{5 + j}", NUM)
    label(ws, "A14", "Tier thresholds (weighted score 0-100)", True)
    label(ws, "A15", "Tier 1 - approach in Wave 1 (score >=)"); inp(ws, "B15", 70, NUM)
    label(ws, "A16", "Tier 2 - approach in Wave 2 (score >=)"); inp(ws, "B16", 50, NUM)
    label(ws, "A17", "Tier 3 - keep informed only (below Tier 2)"); label(ws, "B17", "")
    return ws

# ---------------------------------------------------------------- Market overview / segments
def sheet_market(wb):
    ws = wb.create_sheet("Market_Overview")
    title(ws, "European private debt and alternative credit - market data (sourced)", "Figures as published; where sources conflict both are shown. See Sources tab for URLs.")
    hdr_row(ws, 4, ["Metric", "Value", "Unit", "Period", "Source", "Relevance to Impactus raise"], [52, 16, 12, 16, 46, 70])
    write_rows(ws, 5, MARKET_METRICS)
    ws.freeze_panes = "A5"
    return ws

def sheet_segments(wb):
    ws = wb.create_sheet("Segment_Map")
    title(ws, "Segment map - which part of the credit market fits which Impactus need", "Fit scores 0-5 are analyst judgement (basis stated); pricing ranges are market-observed 2025-2026")
    hdr_row(ws, 4, ["Segment", "Europe size / activity", "Typical pricing", "Typical ticket", "Instruments / hold", "IPL fit (0-5)", "CPEF fit (0-5)", "Why / basis", "Representative firms"], [30, 30, 26, 18, 30, 10, 10, 60, 40])
    write_rows(ws, 5, SEGMENTS, fmts={6: NUM, 7: NUM})
    ws.conditional_formatting.add("F5:G40", ColorScaleRule(start_type="num", start_value=0, start_color="F8CBAD", mid_type="num", mid_value=2.5, mid_color="FFEB9C", end_type="num", end_value=5, end_color="C6EFCE"))
    ws.freeze_panes = "B5"
    return ws

# ---------------------------------------------------------------- Universe (scored)
SCORE_HEADERS = ["Firm", "Category", "HQ / EU office", "Strategy / vehicle", "Fund size / AUM", "Ticket range", "Instruments",
                 "Greece / Cyprus / S. Europe evidence", "Hard-asset / FOAK evidence", "Cannabis stance",
                 "S1 IPL fit", "S2 CPEF fit", "S3 Ticket", "S4 Geography", "S5 Risk appetite", "S6 Warm path", "S7 Speed",
                 "Weighted score", "Tier", "Best use for Impactus", "Recommended channel", "Warm path via", "Source URL"]
SCORE_WIDTHS = [30, 22, 22, 40, 16, 16, 28, 40, 40, 22, 8, 8, 8, 8, 8, 8, 8, 10, 8, 40, 26, 30, 40]

def scored_sheet(wb, name, subtitle, rows, first_col_label="Firm"):
    ws = wb.create_sheet(name)
    title(ws, name.replace("_", " "), subtitle)
    hdrs = list(SCORE_HEADERS); hdrs[0] = first_col_label
    hdr_row(ws, 4, hdrs, SCORE_WIDTHS)
    r = 5
    for row in rows:
        # row: 10 descriptive + 7 scores + best use + channel + warm via + source
        desc = row[:10]; scores = row[10:17]; tail = row[17:]
        for j, v in enumerate(desc, 1):
            c = ws.cell(row=r, column=j, value=v); c.font = BLACK; c.border = BORDER; c.alignment = WRAP
        for j, v in enumerate(scores, 11):
            c = ws.cell(row=r, column=j, value=v); c.font = BLUE_INPUT; c.border = BORDER; c.alignment = CENTER; c.number_format = NUM
        c = ws.cell(row=r, column=18, value=f"=ROUND(SUMPRODUCT(K{r}:Q{r},Scoring_Weights!$E$5:$K$5)/Scoring_Weights!$B$12*20,0)")
        c.font = BOLD; c.border = BORDER; c.alignment = CENTER; c.number_format = NUM
        c = ws.cell(row=r, column=19, value=f'=IF(R{r}>=Scoring_Weights!$B$15,"Tier 1",IF(R{r}>=Scoring_Weights!$B$16,"Tier 2","Tier 3"))')
        c.font = BOLD; c.border = BORDER; c.alignment = CENTER
        for j, v in enumerate(tail, 20):
            c = ws.cell(row=r, column=j, value=v); c.font = BLACK; c.border = BORDER; c.alignment = WRAP
        r += 1
    last = r - 1
    ws.conditional_formatting.add(f"R5:R{last}", ColorScaleRule(start_type="num", start_value=20, start_color="F8CBAD", mid_type="num", mid_value=55, mid_color="FFEB9C", end_type="num", end_value=90, end_color="C6EFCE"))
    ws.conditional_formatting.add(f"S5:S{last}", CellIsRule(operator="equal", formula=['"Tier 1"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(f"S5:S{last}", CellIsRule(operator="equal", formula=['"Tier 2"'], fill=PatternFill("solid", fgColor="FFEB9C")))
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="5", allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"K5:Q{last}")
    ws.freeze_panes = "B5"; ws.auto_filter.ref = f"A4:W{last}"
    return ws, last

# ---------------------------------------------------------------- Contacts
def sheet_contacts(wb):
    ws = wb.create_sheet("Contacts")
    title(ws, "Named contacts - decision-makers, with the channel most likely to get a reply", "Channel rule: warm intro > published email + LinkedIn note same day > LinkedIn InMail with a one-line hook > conference. Emails marked 'pattern - unverified' must be verified (bounce test / mutual contact) before sending. No data was purchased.")
    hdrs = ["Firm", "Person", "Title", "Category", "Target sub-fund", "LinkedIn URL", "Email", "Email status", "Warm path via", "Recommended channel", "Hook / message angle (one line)", "Confidence", "Evidence URL", "Owner", "Status"]
    widths = [28, 24, 34, 20, 12, 44, 34, 18, 28, 26, 60, 11, 44, 8, 14]
    hdr_row(ws, 4, hdrs, widths)
    r = 5
    for row in CONTACTS:
        # row: firm, person, title, category, subfund, linkedin, email, email_status, warm_via, hook, confidence, evidence, owner
        firm, person, ttl, cat, sub, li, em, ems, warm, hook, conf, ev, owner = row
        vals = [firm, person, ttl, cat, sub, li, em, ems, warm, None, hook, conf, ev, owner, "Not started"]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v); c.font = BLACK; c.border = BORDER; c.alignment = WRAP
        ws.cell(row=r, column=10, value=(
            f'=IF(I{r}<>"","Warm intro via "&I{r},IF(H{r}="published","Email + LinkedIn note same day",'
            f'IF(AND(H{r}="pattern - unverified",F{r}<>"not found"),"LinkedIn first; verify email then follow up",'
            f'IF(F{r}<>"not found","LinkedIn InMail / connection note","Firm switchboard or conference"))))'))
        ws.cell(row=r, column=10).font = BOLD; ws.cell(row=r, column=10).border = BORDER; ws.cell(row=r, column=10).alignment = WRAP
        if isinstance(li, str) and li.startswith("http"):
            ws.cell(row=r, column=6).hyperlink = li; ws.cell(row=r, column=6).font = GREEN_LINK
        if isinstance(ev, str) and ev.startswith("http"):
            ws.cell(row=r, column=13).hyperlink = ev; ws.cell(row=r, column=13).font = GREEN_LINK
        r += 1
    last = r - 1
    dv = DataValidation(type="list", formula1='"Not started,Intro requested,Contacted,Replied,Meeting,NDA / data room,Term sheet,Declined,Parked"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"O5:O{last}")
    dv2 = DataValidation(type="list", formula1='"published,pattern - unverified,not found"', allow_blank=True)
    ws.add_data_validation(dv2); dv2.add(f"H5:H{last}")
    ws.freeze_panes = "C5"; ws.auto_filter.ref = f"A4:O{last}"
    return ws, last

# ---------------------------------------------------------------- Outreach plan, conferences, sources
def sheet_outreach(wb):
    ws = wb.create_sheet("Outreach_Plan")
    title(ws, "Outreach plan - waves, channels, cadence, KPIs", "Response-rate assumptions are analyst estimates from placement-agent practice (basis stated); track actuals in the Contacts tab and replace")
    hdr_row(ws, 4, ["Wave", "Timing", "Who (segment)", "Objective", "Channel and cadence", "Material", "KPI / exit criterion", "Owner"], [10, 16, 40, 40, 50, 40, 40, 10])
    r = write_rows(ws, 5, OUTREACH_WAVES)
    r += 1
    label(ws, f"A{r}", "CHANNEL ECONOMICS (assumptions - replace with actuals)", True); r += 1
    hdr_row(ws, r, ["Channel", "Reply rate (assumed)", "Meeting rate (assumed)", "Cost / effort", "When to use", "Basis", "", ""]); r += 1
    chan = [
        ("Warm introduction (Argus / Venero / Sirec / KPMG / PwC / Pyrum / EIB)", "60-80%", "40-60%", "Low cash, high relationship capital", "Always first where a path exists", "Placement-agent convention; Impactus own history (Rodopoulos, Sirec, AstroBank all came via intermediaries)"),
        ("Published work email + same-day LinkedIn connection note", "15-25%", "8-12%", "Low", "Named decision-maker, no warm path", "Cold outreach benchmarks for institutional capital raising (Preqin / placement agents cite 10-20% reply to targeted, personalised notes)"),
        ("LinkedIn InMail / connection note only", "8-15%", "3-6%", "Low", "When email unknown or bounces", "LinkedIn Sales Navigator published InMail benchmarks 10-25% for personalised messages; institutional investors lower"),
        ("Conference meeting request (pre-booked via app)", "30-50%", "20-35%", "High (travel, pass)", "Tier-1 targets clustered at one event", "PDI / SuperReturn organiser data on pre-booked meeting acceptance"),
        ("Placement agent / capital advisor", "n/a", "n/a", "2-5% success fee; retainer risk", "Only for IVP-scale raise with a named target list", "Sergo proposal 5% / 2.5%; KPMG mandate; Venero engaged"),
    ]
    for c in chan:
        write_rows(ws, r, [c + ("", "")]); r += 1
    r += 1
    label(ws, f"A{r}", "MESSAGE ANGLES BY AUDIENCE", True); r += 1
    hdr_row(ws, r, ["Audience", "Lead with (the bear case they will raise, answered)", "Proof point to attach", "Do not say", "", "", "", ""]); r += 1
    for m in MESSAGE_ANGLES:
        write_rows(ws, r, [m + ("", "", "", "")]); r += 1
    ws.freeze_panes = "A5"
    return ws

def sheet_conferences(wb):
    ws = wb.create_sheet("Conferences")
    title(ws, "Where these allocators are in the next twelve months", "Dates as published by organisers at the time of research; verify before booking")
    hdr_row(ws, 4, ["Event", "Dates", "Location", "Audience", "Relevance", "Action", "Source URL"], [40, 18, 18, 40, 40, 40, 44])
    write_rows(ws, 5, CONFERENCES)
    return ws

def sheet_cascade(wb):
    rows = CASCADE_MAP
    ws, last = scored_sheet(wb, "Cascade_Map_2026", "The article's map transcribed: Cascade Debt 'Private ABF Investor Market Map 2026 - Region: Europe', classified by each firm's dedicated ABL/ABF and private credit strategy. Scores are analyst judgement on public evidence.", rows)
    return ws

def sheet_sources(wb):
    ws = wb.create_sheet("Sources")
    title(ws, "Sources", "Every external figure in this workbook traces to one of these; internal figures trace to the Sponsor documents named on the Raise_Parameters tab")
    hdr_row(ws, 4, ["#", "Source", "Publisher / type", "Date", "URL", "Used for"], [6, 50, 26, 12, 70, 40])
    r = 5
    for i, s in enumerate(SOURCES, 1):
        write_rows(ws, r, [(i,) + s]);
        if isinstance(s[3], str) and s[3].startswith("http"):
            ws.cell(row=r, column=5).hyperlink = s[3]; ws.cell(row=r, column=5).font = GREEN_LINK
        r += 1
    return ws

# ---------------------------------------------------------------- Exec summary (formulas)
def sheet_exec(wb, uni_last, casc_last, cont_last):
    ws = wb.create_sheet("Exec_Summary", 1)
    title(ws, "Executive summary - the numbers behind the memo", "All counts are live formulas over the scored tabs")
    ws.column_dimensions["A"].width = 70; ws.column_dimensions["B"].width = 18; ws.column_dimensions["C"].width = 80
    hdr_row(ws, 4, ["Measure", "Value", "Read-across"])
    rows = [
        ("Firms scored in the wider universe (Investor_Universe)", f"=COUNTA(Investor_Universe!A5:A{uni_last})", "Direct lending, special situations, infrastructure credit, venture debt, impact, DFI, cannabis credit, placement, fund finance"),
        ("Firms on the Cascade 2026 Europe ABF map (Cascade_Map_2026)", f"=COUNTA(Cascade_Map_2026!A5:A{casc_last})", "The article's universe, transcribed and scored"),
        ("Tier 1 targets - wider universe", f'=COUNTIF(Investor_Universe!S5:S{uni_last},"Tier 1")', "Approach in Wave 1 (weeks 1-4)"),
        ("Tier 1 targets - Cascade map", f'=COUNTIF(Cascade_Map_2026!S5:S{casc_last},"Tier 1")', "ABF investors that could still do a single hard asset or a receivables/inventory line"),
        ("Tier 2 targets - both tabs", f'=COUNTIF(Investor_Universe!S5:S{uni_last},"Tier 2")+COUNTIF(Cascade_Map_2026!S5:S{casc_last},"Tier 2")', "Wave 2 (weeks 5-10)"),
        ("Named contacts with a warm path", f'=COUNTIF(Contacts!I5:I{cont_last},"<>")', "Use first"),
        ("Named contacts with a published email", f'=COUNTIF(Contacts!H5:H{cont_last},"published")', "Email + LinkedIn note same day"),
        ("Named contacts with LinkedIn only", f'=COUNTIFS(Contacts!F5:F{cont_last},"<>not found",Contacts!H5:H{cont_last},"<>published")', "LinkedIn first"),
        ("Total named contacts", f"=COUNTA(Contacts!B5:B{cont_last})", ""),
        ("IPL: senior debt capacity on model CFADS at 1.5x DSCR (€m)", "=IPL_Debt_Capacity!E16", "vs €77.4m in the Sponsor plan"),
        ("IPL: senior debt capacity at contracted gate fee (€62/t on 11ktpa ~ €15/t on 45ktpa), €1,500/t rCB (€m)", "=PV(IPL_Debt_Capacity!B19,IPL_Debt_Capacity!B20,-(((IPL_Debt_Capacity!B5*IPL_Debt_Capacity!B6*IPL_Debt_Capacity!B7+IPL_Debt_Capacity!B5*IPL_Debt_Capacity!B8*IPL_Debt_Capacity!B9+IPL_Debt_Capacity!B5*IPL_Debt_Capacity!B10*IPL_Debt_Capacity!B11+IPL_Debt_Capacity!B5*15+IPL_Debt_Capacity!B5*IPL_Debt_Capacity!B13)/1000000-IPL_Debt_Capacity!E11)*(1-IPL_Debt_Capacity!B16)-IPL_Debt_Capacity!B5*IPL_Debt_Capacity!B15/1000000)/IPL_Debt_Capacity!B21)", "What a credit committee sizes on today"),
        ("IPL: funding gap vs plan if capex stress applies (€m)", "=IPL_Debt_Capacity!E26", "Negative = shortfall to be filled by private credit or equity"),
        ("IPL: venture-debt all-in cost", "=IPL_Debt_Capacity!E20", "4% cash + 5% PIK + warrants"),
        ("IPL: blended cost of capital", "=IPL_Debt_Capacity!E29", ""),
        ("CPEF: growth-facility headroom at stressed plan (€m)", "=CPEF_Debt_Capacity!E8", "At 2.5x leverage on €3.00/g stressed EBITDA"),
        ("CPEF: all-in cost of a specialist growth facility", "=CPEF_Debt_Capacity!E14", "vs 8% note"),
    ]
    r = 5
    for a, b, c in rows:
        label(ws, f"A{r}", a); fml(ws, f"B{r}", b); label(ws, f"C{r}", c)
        if "€m" in a: ws[f"B{r}"].number_format = EURM
        elif "cost" in a.lower(): ws[f"B{r}"].number_format = PCT
        else: ws[f"B{r}"].number_format = NUM
        r += 1
    return ws

def main():
    wb = Workbook()
    sheet_cover(wb)
    sheet_market(wb)
    sheet_segments(wb)
    sheet_cascade(wb)
    uni_ws, uni_last = scored_sheet(wb, "Investor_Universe", "Wider universe beyond the ABF map: direct lending, special situations, infrastructure and energy-transition credit, venture and growth debt, impact and DFI capital, cannabis-tolerant credit, placement agents and fund-finance providers. Scores are analyst judgement on the public evidence cited.", UNIVERSE)
    cont_ws, cont_last = sheet_contacts(wb)
    sheet_pipeline(wb)
    sheet_ipl_model(wb)
    sheet_cpef_model(wb)
    sheet_outreach(wb)
    sheet_conferences(wb)
    sheet_weights(wb)
    sheet_raise(wb)
    sheet_sources(wb)
    casc_last = 4 + len(CASCADE_MAP)
    sheet_exec(wb, uni_last, casc_last, cont_last)
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = NAVY if ws.title in ("Cover", "Exec_Summary") else ("2E75B6" if ws.title in ("IPL_Debt_Capacity", "CPEF_Debt_Capacity", "Scoring_Weights") else "7F7F7F")
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(OUT)
    print("saved", OUT)

if __name__ == "__main__":
    main()
