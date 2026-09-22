from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.table import WD_TABLE_ALIGNMENT

OUT_DOCX = "/home/user/K/deliverables/HANDOVER_NOTE_EU_Private_Credit_Investor_Map_Sep2026.docx"
OUT_MD = "/home/user/K/deliverables/HANDOVER_NOTE.md"
NAVY = RGBColor(0x1F, 0x2A, 0x44); GREY = RGBColor(0x59, 0x59, 0x59)

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2); s.bottom_margin = Cm(2); s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
for name, size in (("Heading 1", 13), ("Heading 2", 11)):
    h = doc.styles[name]; h.font.name = "Arial"; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = NAVY; h.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

md = []
def H(text, level=1):
    doc.add_heading(text, level=level); md.append(("#" * (level + 1)) + " " + text + "\n")
def P(text, italic=False, size=None, color=None):
    para = doc.add_paragraph(); r = para.add_run(text); r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    para.paragraph_format.space_after = Pt(6); md.append(text + "\n")
def B(text, lead=None):
    para = doc.add_paragraph(style="List Bullet")
    if lead: r = para.add_run(lead + " "); r.bold = True
    para.add_run(text); para.paragraph_format.space_after = Pt(3)
    md.append("- " + ((f"**{lead}** ") if lead else "") + text)
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); tcPr.append(shd)
def T(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(c, "1F2A44"); c.width = Cm(widths[i])
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(8.5); cells[i].width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    md.append("| " + " | ".join(headers) + " |"); md.append("|" + "---|" * len(headers))
    for row in rows: md.append("| " + " | ".join(str(v).replace("|", "/") for v in row) + " |")
    md.append("")

DISC = ("AIFMD Art. 23 disclosure. Impactus RAIF V.C.I.C. Plc is a Registered Alternative Investment Fund (CySEC RAIF16) externally managed by Argus Management Ltd "
        "(CySEC AIFM 44/56/2013). This note is an internal handover document. It is not an offer, solicitation or marketing communication. Interests in the Fund and "
        "its sub-funds are available only to professional and well-informed investors (AIFMD Art. 4(1)(ag), MiFID II, Cyprus AIF Law 124(I)/2018). Targets are gross, "
        "forward-looking and not guaranteed; capital is at risk.")

p0 = doc.add_paragraph(); r = p0.add_run("IMPACTUS PRIVATE EQUITY GROUP - HANDOVER NOTE"); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = GREY
md.append("# Handover note - European private credit investor map and raise analysis (CPEF I Curity, Impactus Pyrolysis)\n")
p1 = doc.add_paragraph(); r = p1.add_run("European private credit investor map and raise analysis - CPEF I Curity and Impactus Pyrolysis (IPL)"); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = NAVY
P("Prepared 22 September 2026 for Konstantinos Michaelides, Partner. Work performed in a Claude Code cloud session on repository michaelidesk/K, branch claude/europe-private-debt-analysis-1x91bg, pull request #2.", italic=True, size=9, color=GREY)
P(DISC, italic=True, size=7.5, color=GREY)

H("1. What was asked and what was delivered")
P("Request: a deep-dive analysis of the market described in Cascade Debt's article 'Private Debt and Alternative Credit Investors in Europe' (the 2026 Europe Private ABF Investor Market Map), an Excel workbook relating that market to the CPEF I Curity and Impactus Pyrolysis raises, and relevant contacts with the best channel to reach each.")
T(["Deliverable", "File", "Status"], [
    ("Analysis memorandum (bear case, market read, models, pre-mortem, decision matrix, position, contact plan, regulatory / tax / liquidity implications, data limits)", "Impactus_EU_Private_Credit_Deep_Dive_Memo_Sep2026.docx", "Complete - 4,500 words, 7 tables, AIFMD Art. 23 disclaimer"),
    ("Working workbook, 15 tabs, live models", "Impactus_EU_Private_Credit_Investor_Map_Sep2026.xlsx", "Complete - 6,909 formula cells verified with a Python formula engine, zero errors; flagged to recalculate on open in Excel"),
    ("Research briefings (market; pyrolysis project finance; cannabis credit; fund-raising routes and Cyprus RAIF rules; Cascade map screen)", "research/*.md (5 files) plus the map image", "Complete - every figure carries a URL; items tagged [K] are analyst knowledge not verified in session"),
    ("Build scripts (regenerate workbook and memo from the data files)", "build/build_xlsx.py, build_memo.py, internal_data.py, external_data.py", "Complete - openpyxl and python-docx"),
    ("This handover note", "HANDOVER_NOTE_EU_Private_Credit_Investor_Map_Sep2026.docx and HANDOVER_NOTE.md", "Complete"),
], [7.0, 6.0, 5.0])

H("2. Where the work lives")
B("GitHub: michaelidesk/K, folder deliverables/, branch claude/europe-private-debt-analysis-1x91bg, pull request #2 (https://github.com/michaelidesk/K/pull/2).", "Repository.")
B("Vertu SharePoint site, Shared Documents / Personal / claude co-work / Impactus / EU Private Credit Investor Map Sep 2026 (this folder).", "SharePoint.")
B("Two files were also sent directly into the chat session on 22 September 2026.", "Chat.")

H("3. Workbook guide")
T(["Tab", "What it holds", "How to use it"], [
    ("Cover", "Purpose, legend, data status, disclaimer", "Read first"),
    ("Exec_Summary", "Live counts of firms, tiers and contacts; model headline outputs", "Recalculates from the other tabs"),
    ("Market_Overview", "24 sourced market data points (fundraising, pricing, defaults, regulation, Greece, cannabis)", "Quote with the source column"),
    ("Segment_Map", "13 credit segments with IPL and CPEF fit scores 0-5", "The argument for which desks to approach"),
    ("Cascade_Map_2026", "The article's 110 map entries transcribed by AUM bucket and scored", "Filter Tier 1 / Tier 2; five illegible logos flagged with candidates"),
    ("Investor_Universe", "78 firms beyond the map: DFIs, Greek banks, special sits, infra / impact debt, venture debt, cannabis credit, placement, fund finance, strategics", "Filter Tier 1; scores are blue inputs you can overwrite"),
    ("Contacts", "68 named entries: person, title, LinkedIn, email, email status, warm path, hook, confidence, owner, status drop-down", "Recommended channel is a formula; status column tracks outreach"),
    ("Existing_Pipeline", "26 relationships already in Impactus records with suggested next action", "Work these before any cold outreach"),
    ("IPL_Debt_Capacity", "PwC 2031 case rebuilt; senior capacity by DSCR; VD cost; sources and gap; two sensitivity grids", "Change blue inputs (gate fee, rCB price, capex stress)"),
    ("CPEF_Debt_Capacity", "Curity stressed plan; leverage-based headroom; cost of a specialist facility; price x volume grid", "Refresh with FY2025 actuals before external use"),
    ("Outreach_Plan", "Waves 0-3 with timing, channel cadence, KPIs; channel economics; message angles per audience", "Operating plan for the next six months"),
    ("Conferences", "14 events Nov 2026 to Jun 2027 with relevance and action", ""),
    ("Scoring_Weights", "Seven criteria weights and tier thresholds", "Change here and every score recalculates"),
    ("Raise_Parameters", "24 internal reference points from the Sponsor's own September 2026 documents", ""),
    ("Sources", "48 sources with URLs", ""),
], [3.6, 8.0, 6.4])

H("4. Key findings to carry forward")
B("Roughly 80% of the Cascade map (CLO, ABS, SRT, fintech warehouse, NPL, secondaries) has no product for a single first-of-a-kind plant or an early-revenue cannabis producer. The relevant fifth is the special-situations, real-asset and energy-credit desks.", "Fit.")
B("Senior capacity on the PwC case computes to €78.2m against Credia's €77.4m; on the contracted gate fee (c.€15/t across 45ktpa) it is €64.0m. Each €10/t of gate fee on full throughput is worth c.€1.7m of senior debt. On contracted revenue plus a €150m plant capex the gap is c.€45m, not €35m.", "IPL model.")
B("The €35m venture-debt tranche at 10% blended is below where FOAK junior capital clears (14% to 18% all-in with 10% to 20% warrant coverage). The EIB declined on exactly this point on 30 April 2026 and reopened the file in September via Sirec.", "Pricing.")
B("Every financed European tyre-pyrolysis project (Circtec, Bolder, Infiniteria, Wastefront) was funded by grant plus strategic equity plus offtaker-linked debt; none by fund senior debt. Offtakers (BASF, Continental / Pirelli via Pyrum, Archirodon) belong in the stack first.", "Precedent.")
B("Curity: stressed-plan headroom for a growth facility is €7.5m at c.14% all-in; the realistic lenders in 2026 are strategics and specialist cannabis funds; European growth-debt houses need twelve months of audited sales.", "CPEF model.")
B("Recommended option B in the memo: offtakers first, EIB with equity-like reward, HIIF staged, Ecoelastika amendment before Credia committee, one competitive hybrid process (Hayfin Special Opps, AB CarVal clean energy, Sienna) via Venero, and the €5m IPL round sold as a sponsor bridge to the Cyprus network.", "Position.")

H("5. Contacts - status and rules")
B("68 named entries: 34 from Impactus' own records (24 with published emails, the rest warm via a named introducer) and 34 external targets.")
B("External names are marked 'verify' and carry a LinkedIn search link, not a fabricated profile URL, because web access and Apollo enrichment were unavailable (see section 7). Verify each on LinkedIn before sending.")
B("Channel rule embedded as a formula: warm introduction first; then published email plus a LinkedIn connection note the same day; then LinkedIn InMail; then conference. Stop after three touches.")
B("Live threads with emails on file: EIB (Casorati, Mitrakas, Gaudet), HIIF (Vlachos, Fragos, Giannakopoulos), Credia (Mitrou, Masouras, Iliopoulos), Sirec (Papageorgantas, Alissandratos), Venero, PwC Greece Deals, Fordoers, Muhanna, Axia, J. Safra Sarasin, IDN.")

H("6. Immediate actions with owners (from the memo, Section 7.1 and Outreach_Plan Wave 0)")
T(["#", "Action", "Owner", "By"], [
    ("1", "Deliver the three EIB committee items: Credia CP explainer, Continental supplier-approval letter (via Pyrum), Sokolov comparison; propose warrant coverage 15% to 20% or a conversion right", "KM / Sirec", "This week"),
    ("2", "Re-cut the HIIF €35m into a development tranche now and a construction tranche at FID; agree the market-operator test basis", "KM", "Before the PMO reset is communicated"),
    ("3", "Sign Ecoelastika amendment No.2 extending the gate fee to full throughput", "KM", "Before Credia credit committee"),
    ("4", "Firm EPC price and TÜV Nord TDD summary to Credia; ask for committee calendar and CP list in writing; run Piraeus or Eurobank PF in parallel via PwC Deals", "Venero / KM", "31 October 2026"),
    ("5", "Offtaker prepayment asks: BASF (oil), Continental / Pirelli via Pyrum (rCB Line A), Archirodon deferred EPC payment", "KM", "October 2026"),
    ("6", "Venero success-fee mandate to open Hayfin Special Opps, AB CarVal clean-energy credit, Sienna Private Credit; Rivage ECDS and Eiffel (grant bridge) as the impact pair", "KM / Venero", "October 2026"),
    ("7", "PMO reset: require end-date, budget and scope boundary before it reaches CINEA, HIIF or the EIB; re-baseline the grant milestone schedule with CINEA in the October amendment", "KM / PwC", "End October 2026"),
    ("8", "Close the EY IFRS 13 fair-value query for the Doctors' Provident Fund; produce a coupon-coverage bridge for the PE Fund before any further provident-fund presentations", "KM / MI", "October 2026"),
    ("9", "CPEF: re-open ELPEN; approach Cantourage or Curaleaf International for offtake; test Chicago Atlantic (John Mazarakis) and Artemis Growth Partners for a 2027 facility", "AM / KM", "Q4 2026"),
    ("10", "Verify the 34 external contact rows on LinkedIn; enrich emails once Apollo is on a paid plan", "MI / LP", "Rolling"),
], [0.8, 10.0, 3.2, 4.0])

H("7. Limitations and what could not be verified")
B("The Cascade page, all firm websites, LinkedIn and the press wires were blocked by the sandbox network policy; the session's web-search quota ran out after about 200 queries. Every external figure comes from search-engine extracts of the cited page, not a full read. Re-verify before quoting externally.", "Web access.")
B("Apollo.io is on a free plan that blocks API search and enrichment; no contact data was purchased.", "Apollo.")
B("LibreOffice Calc is not installed in the build environment; the workbook was verified with a Python formula engine (6,909 cells, zero errors) and is flagged to recalculate on open in Excel.", "Recalculation.")
B("Five logos on the map are illegible; candidates (Fasanara, Fedaia, Avenue Capital) are flagged as unverified. Ask Cascade Debt for the underlying list.", "Map.")
B("Curity figures rest on the FY2023 Meerkat plan; the FY2025 valuation and audited accounts should replace them before any lender sees the CPEF tab.", "Curity data.")
B("The capex benchmark in the PMO minutes ('100 vs 250') is recorded without units; plant capex is unresolved between €118.5m and €150m until the firm EPC offer of 31 October 2026.", "IPL data.")

H("8. Internal sources used (all in Impactus SharePoint / Outlook / Fathom)")
B("Thermolysi Project - Company Snapshot Sep 2026 v2 and Indicative Term Sheet 11-09-2026 (Thermo Lysi fundraise Sep 2026 folder).")
B("R11 - PMO Change of Strategy - MoM & Assessment 03-09-2026.")
B("Impactus Investor Pack Q2 2026 v4; 8% CPEF Curity Note term sheet and email of 08-05-2026; Meerkat Curity valuation FY2023 (Appendix IV projections).")
B("Outlook threads: Credia <> Impactus CINEA bank letter (Jul-Aug 2026); Thermo Lysis Financial Model with EIB Cleantech (Feb-Apr 2026); EIB responses - follow up (15-09-2026); Impactus Group - HIIF (Jul-Aug 2026); Re: Impactus Pyrolisis with Muhanna / EY (17-09-2026); investor tracker LIST OF INVESTORS.xlsx (tabs to 03-08-2026); Cooperation Alpha Bank; IVP to J. Safra Sarasin.")
B("Fathom calls: Sergo Vashakmadze 23-06-2025; IG Funds / Oak Chase Global 20-03-2026; Impactus / Pyrum 05-02-2026.")

H("9. How to update the deliverables")
B("Edit deliverables/build/external_data.py (market, segments, universe, contacts, map, conferences, sources) or internal_data.py (raise parameters, pipeline).")
B("Run python deliverables/build/build_xlsx.py then python deliverables/build/build_memo.py; both scripts write into deliverables/.")
B("Requirements: Python 3.11, openpyxl, python-docx. Optional: LibreOffice Calc for a full recalculation; otherwise Excel recalculates on open.")
B("Commit to the branch and push; pull request #2 updates automatically.")

P("")
P(DISC, italic=True, size=7.5, color=GREY)
doc.save(OUT_DOCX)
open(OUT_MD, "w").write("\n".join(md))
print("saved", OUT_DOCX, OUT_MD)
