from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/home/user/K/deliverables/Impactus_EU_Private_Credit_Deep_Dive_Memo_Sep2026.docx"
NAVY = RGBColor(0x1F, 0x2A, 0x44); GREY = RGBColor(0x59, 0x59, 0x59)

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2); s.bottom_margin = Cm(2); s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
    h = doc.styles[name]; h.font.name = "Arial"; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = NAVY
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

def p(text, bold=False, italic=False, size=None, color=None, align=None, space_after=6):
    para = doc.add_paragraph(); r = para.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if align: para.alignment = align
    para.paragraph_format.space_after = Pt(space_after)
    return para

def bullet(text, bold_lead=None):
    para = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        r = para.add_run(bold_lead + " "); r.bold = True
    para.add_run(text); para.paragraph_format.space_after = Pt(3)
    return para

def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill); tcPr.append(shd)

def table(headers, rows, widths_cm, font=8.5):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(font); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "1F2A44"); c.width = Cm(widths_cm[i])
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(font); cells[i].width = Cm(widths_cm[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

DISCLAIMER = ("AIFMD Art. 23 disclosure. Impactus RAIF V.C.I.C. Plc is a Registered Alternative Investment Fund (CySEC RAIF16) externally managed by "
              "Argus Management Ltd (CySEC AIFM 44/56/2013). This memorandum is an internal analytical document prepared for the Sponsor. It is not an "
              "offer, solicitation or marketing communication in any jurisdiction. Interests in the Fund and its sub-funds (Impactus Private Equity Fund, "
              "CPEF I Curity, Impactus Pyrolysis / IVP) are available only to professional and well-informed investors within the meaning of AIFMD "
              "Art. 4(1)(ag), MiFID II and the Cyprus Alternative Investment Funds Law 124(I)/2018. Prospective investors must read the AIFMD Art. 23 "
              "disclosure document, the relevant offering supplement and term sheet before any investment decision. Targets are gross, forward-looking "
              "and not guaranteed; capital is at risk. Third-party names are used for identification only and do not imply endorsement.")

# ------------------------------------------------------------------ Title block
p("IMPACTUS PRIVATE EQUITY GROUP", bold=True, size=9, color=GREY)
p("European Private Debt and Alternative Credit Investors", bold=True, size=18, color=NAVY, space_after=2)
p("Relevance to the CPEF I Curity and Impactus Pyrolysis (IPL) raises - deep-dive analysis, investor map and contact plan", size=11, color=GREY, space_after=2)
p("Internal analysis memorandum - 22 September 2026 - prepared for K. Michaelides, Partner. Companion workbook: Impactus_EU_Private_Credit_Investor_Map_Sep2026.xlsx", italic=True, size=9, color=GREY, space_after=10)
p(DISCLAIMER, italic=True, size=7.5, color=GREY, space_after=12)

# ------------------------------------------------------------------ 1. Synthesis
doc.add_heading("1. The core truth", level=1)
p("The European private credit market is large, liquid and rotating towards Europe, and almost none of it is built for what Impactus needs. "
  "The Cascade article maps roughly one hundred asset-based finance (ABF) investors; ABF lends against granular pools of receivables, loans and leases "
  "with observable performance data. Thermo Lysi is a single first-of-a-kind (FOAK) plant with no operating history, and Curity Pharma is an early-revenue "
  "producer in a sector most institutional lenders screen out. Of the article's universe, roughly one fifth - the special-situations, real-asset and "
  "energy-credit desks - could in principle write a €20m to €40m secured hybrid behind a Greek bank senior facility, and none could write the €5m equity "
  "ticket currently on offer. The money that has actually financed European tyre pyrolysis in 2024 to 2026 came from grants, strategic equity and "
  "offtaker-linked debt, not from credit funds. The practical conclusion is that the IPL capital plan should be re-sequenced around the offtakers, the "
  "reopened EIB venture-debt file and the Hellenic Innovation and Infrastructure Fund (HIIF), with one special-situations lender run as competitive tension; "
  "and that CPEF should not spend effort on institutional credit until twelve months of audited German sales exist.")

# ------------------------------------------------------------------ 2. Bear case
doc.add_heading("2. The bear case first - argued in the counterparties' own words", level=1)
p("Each objection below is the one the named actor has already raised, or will raise in week one of diligence. The position in Section 7 is built to survive them, not to avoid them.")

doc.add_heading("2.1 Alberto Casorati, EIB Cleantech Equity and Growth Capital (30 April 2026)", level=2)
p("\"While the EIB venture debt with its quasi-equity characteristics is designed to take such risks (often in the context of First Of A Kind projects) along with the equity investors, "
  "we want to avoid deploying our risk capital as a mezzanine exposing us to equity risk but not to its reward.\" That is the whole case against the €35m venture-debt tranche as "
  "currently priced: 4% cash plus 5% PIK plus a warrant, a 10% blended return, ranking ahead of €21m of shareholder capital but below the senior facility, on a plant whose licensor "
  "(Pyrum) reported FY2025 revenue of €4.1m and delayed break-even. The EIB reopened the file in September through Sirec and is asking three concrete questions. The tranche will "
  "not clear at 10% blended with a token warrant. Junior capital for a FOAK industrial asset clears at 14% to 18% all-in with 10% to 20% warrant coverage.", italic=False)

doc.add_heading("2.2 Credia Bank credit committee", level=2)
p("The senior facility is sized at 1.5x DSCR on the PwC model, which assumes a €100/t gate fee across 45ktpa and €1,500/t for recovered carbon black. The contracted position is "
  "€62/t on 11ktpa (about €15/t across the full plant) and the rCB price is a CIF Dillingen reference that Impactus itself haircuts to €1,350/t for logistics. The committee will "
  "size on contracted cash flows only. On the model, senior capacity at 6.5% over ten years is €78.2m, which reconciles to the €77.4m in the Sponsor plan; on the contracted gate "
  "fee it is €64.0m. Every €10/t of gate fee on the full throughput is worth about €1.7m of senior debt. The Ecoelastika amendment is therefore worth more than any lender meeting.")

doc.add_heading("2.3 The PMO (Fordoers, 2 September 2026)", level=2)
p("Their own deck says December FID is not feasible: capex intensity unresolved between €120m and €150m for the plant, textile content at 24% to 25% against a 12% specification, "
  "a quarter of the components carrying a design or legality gap, and revenue adequacy resting on the gate fee. Munich Re withdrew its technology performance insurance product in "
  "February 2026. A credit investor who reads the PMO minutes before they are re-baselined with an end-date and a budget will price the reset as an open-ended delay. The reset must "
  "reach CINEA, HIIF and the EIB as a dated milestone schedule, not as a strategy narrative.")

doc.add_heading("2.4 The LP due-diligence lead at a Cyprus provident fund, and EY as auditor", level=2)
p("EY has already asked, on behalf of the Doctors' Provident Fund, for audited accounts, the valuation basis and the workings behind a fair-value uplift; the latest audited accounts "
  "of Impactus Pyrolysis Ltd are FY2023 with a qualified opinion. The Investor Pack states an 8% net coupon paid every term since 2022 while both platform assets are pre-revenue, and "
  "the Sponsor's own August 2025 note identified about €3m of new money needed for 2027 to 2028 operating costs and coupon. A DD lead will ask whether the coupon is serviced from asset "
  "cash flows or from subscriptions. Until a cash bridge answers that question in writing, every institutional conversation about the PE Fund is exposed.")

doc.add_heading("2.5 A special-situations portfolio manager", level=2)
p("No European tyre-pyrolysis project has closed non-recourse senior debt from a bank or an infrastructure-debt fund. Circtec (€150m) was funded by Novo Holdings and A.P. Moller "
  "equity, Dutch grants and a €12.5m bp offtaker loan; Bolder Antwerp by Tiger Infrastructure equity and a €32m Innovation Fund grant; Infiniteria by Antin equity, and it is now in "
  "reorganisation and arbitration. Plastic Energy's UK entities went into administration in April 2026, Viridor proposed closing Quantafuel in May 2026, Brightmark defaulted on "
  "US$172.5m of green bonds. The PM's opening position is that this asset class has a pattern of technology and offtake failure, and that the correct instrument is equity or "
  "offtaker capital, not fund debt. That is why the price of any hybrid will be high and why the offtakers must be in the stack first.")

doc.add_heading("2.6 A specialist cannabis lender's credit officer", level=2)
p("German pharmacy flower prices fell from €8.33/g in January 2025 to €4.52/g in March 2026; statutory health-insurance reimbursement of flower ended on 30 July 2026; imports rose "
  "2.8x to 201 tonnes in 2025, so volume is not the problem, price is. Every European distressed exit recovered a fraction of invested capital (Clever Leaves' Portuguese assets sold "
  "for €2.5m; RPK / Holigen for US$2.0m). Advanced Flower Capital carries 35.7% of its book on non-accrual. No specialist lender has a documented European loan. The officer will "
  "underwrite at the BfArM tender reference of €2.30/g, not at Curity's €4.85/g base case, and will lend only against hard security and contracted German offtake.")

# ------------------------------------------------------------------ 3. Market
doc.add_heading("3. What the article's market actually is", level=1)
p("Cascade Debt's page is the '2026 Europe Private ABF Investor Market Map': about a hundred firms classified by each firm's dedicated ABL/ABF and private-credit strategy, "
  "bucketed by AUM (under US$5b, US$5b to US$20b, US$20b to US$100b, over US$100b). The page itself is blocked from this environment; the map image you supplied has been transcribed "
  "firm by firm into the Cascade_Map_2026 tab and scored. Five logos are illegible; candidate identifications (Fasanara, Fedaia, Avenue Capital) are flagged as unverified.")
table(["Metric", "Figure", "Source", "Why it matters here"], [
    ("Europe private debt fundraising 2025", "US$79.4b (PitchBook); US$69.9b (S&P / With Intelligence); second-highest year", "PitchBook; S&P", "Capital is available; fit is the constraint"),
    ("Europe share of global private credit fundraising", "23% (2024) to 46% (2025)", "Alternative Credit Investor, 16-12-2025", "The window for European stories is open"),
    ("Direct lending share of European capital raised", "61.5% (9M 2025)", "Preqin via ACI", "Most of the money is sponsor-backed corporate lending, not project or growth credit"),
    ("Concentration", "93.1% of 2025 fundraising to managers on Fund IV or later", "PitchBook 2025 annual report", "A first-time sub-€100m GP will not raise from these allocators directly"),
    ("Direct lending deal volume Q2 2026", "€28.4b / 296 deals, down 25% year on year", "ACI, 25-08-2026", "Lenders are competing for assets; pricing power is with borrowers who fit"),
    ("Unitranche pricing, mid-market", "Euribor + 525 to 625bps; 5.0x to 6.5x EBITDA", "Valuation Research Corp Q2 2026", "Benchmark for a corporate credit; Thermo Lysi is not that credit"),
    ("Special-situations fundraising H1 2026", "Alchemy, Metric, Tikehau each closed €1b+; pipeline c.£100b", "With Intelligence; ACI", "The mandate for PIK, warrants and FOAK sits here"),
    ("European infrastructure debt 2025", "Record £68b financed; Infranity €3.2b, Eiffel €1.2b, Rivage €700m target", "Aviva; ACI; Infrastructure Investor", "Enhanced-return and climate-debt sleeves are the only infra-debt entry for a FOAK plant"),
    ("ECB deposit rate", "2.00% to 2.50% between June and September 2026", "ECB", "Every senior term sheet quoted in H1 2026 is 25 to 50bps stale"),
    ("Private credit stress signals", "Proskauer default 2.51% (Q2 2026); Lincoln 'bad PIK' 6.4%; Fitch US record 6.0%", "Proskauer; Lincoln; Fitch", "Committees are tightening on PIK - the 5% PIK in the VD tranche will be read as a warning sign"),
    ("AIFMD II and Cyprus", "Transposition due 16-04-2026; Cyprus not fully transposed; RAIFs may not run loan-originating strategies", "Jones Day; Chambers; LCK", "Note programmes must stay at holdco level, outside the RAIF's perimeter"),
    ("EIB Group in Greece 2025", "Record €3.5b; venture debt €17.5m to Joltie", "EIB press 2026-041", "EIB venture debt exists in Greece at the €15m to €20m scale"),
], [4.2, 5.2, 3.2, 5.2])

doc.add_heading("3.1 Which segment fits which need", level=2)
table(["Segment", "IPL fit (0-5)", "CPEF fit (0-5)", "Verdict"], [
    ("ABF (the Cascade universe)", "1", "2", "Granular pools only; single FOAK plant does not fit; Curity receivables line possible in 2027"),
    ("Sponsor-backed direct lending", "0", "1", "Needs audited EBITDA and usually a PE sponsor; exclude"),
    ("Lower-mid-market / SME credit (Southern Europe)", "2", "3", "Only core-credit segment writing €3m to €30m; Kartesia, Capza, Muzinich; Greek AIFs (Elikonos, Halcyon)"),
    ("Special situations / opportunistic", "4", "2", "Natural home for the €35m hybrid at 14% to 18% all-in; cannabis case by case"),
    ("Infrastructure and energy-transition debt", "3", "0", "Enhanced-return sleeves only (Infranity ERDF, Rivage High Yield, Eiffel bridge, Sienna, AB CarVal clean energy)"),
    ("Venture / growth debt", "2", "3", "Right ticket, wrong stage for IPL; Curity at 12 months of audited sales; Kreos (Constantinides) is the Greek door"),
    ("Development finance", "4", "1", "EIB VD reopened; EBRD Greece closed Dec 2025; HDB / RRF windows closed Aug 2026"),
    ("Impact / climate credit", "4", "0", "Best mandate fit: Polestar (Xycle lender), Rivage ECDS, Blue Earth, Triodos"),
    ("Strategic / offtaker capital", "5", "3", "The pattern in every financed pyrolysis project; BASF, Continental / Pirelli via Pyrum, Archirodon; ELPEN and German importers for Curity"),
    ("Specialist cannabis credit", "0", "4", "Chicago Atlantic (Mazarakis), AFC, Artemis Fund VII (75% Europe); 13% to 16%; no European loan yet"),
    ("Fund finance (NAV / capital-call)", "1", "0", "Every institutional lender is above scale; UK lenders apply POCA to cannabis; only Eurobank Cyprus or BoC could write a small line"),
    ("Cyprus / Greek institutional and family capital", "3", "3", "Where the money has actually come from; scale via Argus, Muhanna, Axia and an IORP-compliant note"),
], [5.2, 2.0, 2.0, 8.6])

# ------------------------------------------------------------------ 4. IPL
doc.add_heading("4. IPL / Thermo Lysi - what the credit market can carry", level=1)
p("The IPL_Debt_Capacity tab reproduces the PwC 2031 case (44,776t; rCB 30.3% at €1,500/t; oil 20% at €845/t; steel 24.5% at €250/t; gate fee €100/t; carbon €25/t; cash opex €317/t) "
  "and sizes senior debt as the present value of CFADS divided by DSCR over the repayment tenor at the facility rate. The results below are the model's outputs; the inputs are live.")
table(["Measure", "Value", "Read-across"], [
    ("2031 EBITDA on the PwC case", "€22.1m (PwC: €21.9m)", "Model reconciles to the Sponsor's figures"),
    ("Senior capacity at 1.5x DSCR, 6.5%, 10 years", "€78.2m", "Matches Credia's €77.4m - the plan is internally consistent on its own assumptions"),
    ("Senior capacity on the contracted gate fee (c.€15/t across 45ktpa)", "€64.0m", "€14m short of plan - this is what a credit committee sizes on today"),
    ("Sensitivity", "Each €10/t of gate fee on full throughput = c.€0.45m EBITDA = c.€1.7m senior capacity", "The Ecoelastika amendment is the highest-leverage line in the pack"),
    ("Sources check on the Sponsor plan", "€176.1m sources against €176m required", "Balanced only if Credia, HIIF / VD and the VAT facility all close at plan size"),
    ("Gap if plant capex is €150m rather than €118.5m", "€31.4m", "PMO: '€120m or €150m' unresolved"),
    ("Gap if senior is sized on contracted revenue and capex stress applies", "c.€45m", "This is the true size of the private-credit ask, not €35m"),
    ("Venture-debt all-in cost as offered", "9.4% (4% cash + 5% PIK + 3% warrant PV over 7 years)", "Below where junior FOAK capital clears; expect 14% to 18%"),
    ("Blended cost of capital (senior, VD, shareholder loans, equity at 22%)", "9.2%", ""),
    ("Interest cover, EBITDA to cash interest (senior plus VD)", "3.4x", "Adequate at run-rate; the risk is construction and ramp-up, not run-rate coverage"),
], [6.2, 5.0, 6.6])
p("DD challenge. The model EBITDA of €22m rests on two prices that are not contracted at that level today. The venture-debt or private-credit tranche is not a cheaper alternative "
  "to equity; it is the price of the gap between contracted and modelled revenue. Presenting it at 10% blended to a credit investor invites the EIB's April answer again.", bold=False, italic=True)

# ------------------------------------------------------------------ 5. CPEF
doc.add_heading("5. CPEF I Curity - the honest price of growth debt", level=1)
p("The CPEF_Debt_Capacity tab uses the Meerkat stressed plan (3.5m grams at €3.00/g, 50% EBITDA margin) and a 2.5x leverage ceiling, which is where specialist cannabis lenders sit.")
table(["Measure", "Value", "Read-across"], [
    ("Stressed revenue / EBITDA at capacity", "€10.5m / €5.2m", "Meerkat FY2023 plan; refresh with FY2025 actuals before external use"),
    ("Debt capacity at 2.5x", "€13.1m", ""),
    ("Headroom after existing bank debt (€5.6m)", "€7.5m", "A €3m to €10m growth facility is feasible on the stressed case at capacity, not today"),
    ("Cost of a specialist facility", "13% cash + 4% fees and warrants = c.14% all-in", "600bps above the 8% note; the note is priced by relationship, not by the market"),
    ("Interest cover at that structure", "4.0x", ""),
    ("Underwriting price a lender will use", "€2.30/g (BfArM reference), trending lower", "At €2.50/g and 2.5m grams the headroom falls to zero"),
], [6.2, 5.0, 6.6])
p("DD challenge. Curity is early-revenue (first exports 2026; Galenica distribution from October 2026). The realistic lenders in 2026 are strategics (ELPEN, German importers) and "
  "specialist cannabis funds; European growth-debt houses (Kreos, Claret, Atempo) have no stated cannabis exclusion but require twelve months of revenue. The 8% note is therefore "
  "the right instrument for now, and it should be sold as a relationship product, not as a market-priced credit.", italic=True)

# ------------------------------------------------------------------ 6. Pre-mortem
doc.add_heading("6. Pre-mortem", level=1)
table(["Scenario", "What happens", "Capital impact", "Probability anchor"], [
    ("Base case", "Firm EPC price lands at €120m to €125m for the plant; Ecoelastika amendment extends the gate fee to full throughput at €60/t to €70/t; Credia approves c.€65m to €70m; EIB VD re-cut at €20m to €25m with real warrant coverage; HIIF commits in two stages; FID slips to Q2 2027",
     "Sponsor-level gap of €10m to €15m to be filled by offtaker prepayments (BASF, Continental via Pyrum) and the €5m IPL round; VD cost rises to c.14% all-in", "What I would bet on at fair odds - the PMO evidence supports a six-month slip, not a cancellation"),
    ("Bear case", "EPC firm price comes in at €150m; textile and configuration gaps push TDD sign-off to mid-2027; HIIF declines on the market-operator test; EIB declines a second time; Credia sizes on contracted revenue only",
     "Funding gap of €40m to €50m; the project cannot reach FID without an industrial partner taking equity; IPL's €50m pre-money valuation is not defensible and the €5m round stalls", "One in three - the PMO has already told you three of the four blockers are technical"),
    ("Black swan", "Pyrum enters restructuring (FY2025 revenue miss, breakeven delayed, €13m of a €21m rights issue placed) or Sokolov / Perl-Besch ramp-up fails publicly; Continental pauses rCB integration; CINEA reassesses the grant on entry-into-operation risk",
     "Technology licensor and offtake chain fail together; Thermo Lysi becomes a permitted site with a grant option; IPL carrying value (€21.5m Meerkat) written down materially; CPEF unaffected", "Low probability, severe consequence - the licensor is a sub-€15m revenue company carrying the whole platform"),
], [2.4, 6.4, 5.0, 4.0])

# ------------------------------------------------------------------ 7. Decision matrix and position
doc.add_heading("7. Decision matrix and position", level=1)
table(["", "A. Current plan", "B. Strategic-first stack (recommended)", "C. Special-sits hybrid replaces VD", "D. Development-partner reset and industrial sale"], [
    ("Core thesis", "Credia senior + €35m VD at 10% + €21m shareholders + €5m IPL round; FID Dec 2026", "Offtaker prepayments (BASF oil, Continental / Pirelli rCB via Pyrum, Archirodon deferred EPC) first; EIB VD re-cut with equity-like reward; HIIF staged; Credia sized on contracted revenue; one special-sits lender run for tension",
     "Replace the VD tranche with a €25m to €40m hybrid from Hayfin Special Opps, Blantyre, Incus, Cross Ocean or AB CarVal at 14% to 18% with warrants", "Accept the PMO diagnosis: hold and de-risk for 12 to 18 months, then sell a permitted, contracted, grant-anchored site to an industrial (Continental, Pirelli, VTTI, Circtec, Bolder)"),
    ("Key risks / mitigants", "VD does not clear at 10%; Credia sizes down; FID slips - no mitigant inside the plan", "Offtakers may wait for market proof - mitigant is the VTTI Antwerp 90ktpa threat and a floor-price prepayment; EIB may decline again - mitigant is the reward fix",
     "Price dilutes equity returns to c.15% IRR; PIK read as distress - mitigant is cash-pay step-up at COD", "Value of goodwill unproven; grant timetable (Art. 39 amendment) must be re-baselined with CINEA; PMO extends its own mandate - mitigant is a costed roadmap with an end date"),
    ("Capital impact", "Gap €0 on paper; €31m to €45m on contracted revenue and capex stress", "Closes €15m to €25m of the gap at below-market cost; VD cost rises to c.14%; blended cost of capital c.10%", "Closes the tranche; blended cost of capital c.11%; equity IRR falls from 22% target to c.15% to 17%", "No construction capital raised; holding cost €3m to €5m over 18 months; exit value depends on industrial appetite (Antin, Tiger, Novo precedents)"),
    ("Strategic alignment", "IVP thesis intact only if FID holds", "Keeps IVP thesis and Greek exclusivity; converts the €5m IPL round into a bridge with a named use of proceeds", "Keeps FID path; weakens IVP economics", "Converts IVP from operator to developer - a different fund, a different LP base"),
], [2.6, 3.6, 4.2, 3.6, 3.6], font=8)

doc.add_heading("7.1 Position", level=2)
p("Option B, with C prepared as the fallback and D kept as the honest alternative if the firm EPC price exceeds €140m. Specifically:")
bullet("Sequence the stack around the counterparties who already have skin in the outcome: BASF (€16m equity and a €25m conditional loan into Pyrum), Continental (advance payment precedent), Pirelli (July 2026 supply agreement), Archirodon (interest in operations). Ask each for a prepayment or deferred payment against Line A, on the bp / Circtec model. This is the only capital in Europe that has actually financed tyre pyrolysis.", "Offtakers before lenders.")
bullet("Answer the April objection rather than re-arguing it: raise warrant coverage to 15% to 20% or add a conversion right, present a P90 case on contracted revenue, deliver the Continental supplier-approval letter and the Sokolov comparison this week.", "EIB: give the reward.")
bullet("Re-cut the €35m into a development tranche now and a construction tranche at FID; agree the market-operator test basis (pari passu with Sirec's executed €34m post-money round) before the PMO reset reaches them.", "HIIF: stage it.")
bullet("Sign Ecoelastika amendment No.2 to full throughput before the Credia credit committee; the model shows €1.7m of senior capacity per €10/t. Run Piraeus or Eurobank in parallel through PwC Deals for tension.", "Credia: fix the gate fee first.")
bullet("Mandate Venero, on success fee, to open exactly three: Hayfin Special Opportunities, AB CarVal clean-energy credit and Sienna Private Credit (energy debt at €5m to €50m), with Rivage ECDS and Eiffel (grant-receivable bridge) as the impact-debt pair. Do not approach the wider Cascade universe; 80% of it has no product for this risk.", "One competitive hybrid process, not a broad one.")
bullet("Sell it as what it is: a sponsor-level bridge to FID with named uses (ECI fees €0.86m, engineering, PMO, TDD, permitting), to the Cyprus network that has already funded €5.5m (SEK, Doctors' Fund, CSE fund, AstroBank, Rodopoulos) plus ETYK and PEO. Do not present a €50m pre-money valuation to institutional credit; they will not underwrite it and it costs credibility.", "The €5m IPL round.")
bullet("Keep the 8% note as the instrument; add ELPEN and one German importer (Cantourage or Curaleaf International) as strategic co-investors; test Chicago Atlantic through John Mazarakis and Artemis Growth Partners for a 2027 growth facility once Galenica sales are audited. Do not approach mainstream direct lenders or any UK lender (POCA).", "CPEF.")
bullet("Before any institutional LP conversation about the PE Fund, produce a coupon-coverage bridge and close the EY fair-value query. These two items decide whether the Cyprus provident-fund channel scales or stalls.", "Fund level.")

# ------------------------------------------------------------------ 8. Contacts
doc.add_heading("8. Who to contact and how", level=1)
p("The Contacts tab holds 68 named entries: 34 drawn from Impactus' own records (24 with published emails, the rest warm via a named introducer) and 34 external targets. "
  "Apollo.io enrichment and web search were unavailable in this session, so external names are drawn from analyst knowledge, marked 'verify', and carry a LinkedIn search link "
  "rather than a fabricated profile URL. Channel rule: warm introduction, then published email plus a LinkedIn note the same day, then LinkedIn InMail, then conference. "
  "Cold email to institutional investors goes 95% unanswered; a warm introduction raises meeting probability roughly tenfold (sources in the workbook).")
table(["Priority", "Firm / person", "Channel", "The one-line hook"], [
    ("1", "EIB - Alberto Casorati, Stephan Mitrakas (a.casorati@eib.org, s.mitrakas@eib.org)", "Existing thread via Sirec; call then email", "Equity-like reward added to the VD; three committee items delivered"),
    ("2", "HIIF - Alexandros Vlachos, Stelios Fragos, Andreas Giannakopoulos (@hiif.gr)", "Existing thread; working session", "Staged commitment; market-operator test on the Sirec round"),
    ("3", "Credia - Evangelos Mitrou (mitrou.evangelos@crediabank.com)", "Venero; weekly", "Firm EPC price plus Ecoelastika amendment before committee"),
    ("4", "Ecoelastika board", "Direct", "Amendment No.2 to full throughput at a fixed fee"),
    ("5", "Pyrum - Pascal Klein; BASF ChemCycling; Continental via Pyrum", "Direct (KM)", "Supplier-approval letter; prepayment against Line A"),
    ("6", "Hayfin Special Opportunities; AB CarVal clean-energy credit; Sienna Private Credit", "Venero mandate; LinkedIn first", "€25m to €40m hybrid behind Credia, senior to €21m of shareholder capital"),
    ("7", "Rivage ECDS; Eiffel Energy Transition; Polestar Capital", "LinkedIn first", "Grant-receivable bridge; Xycle covenant template"),
    ("8", "Tiger Infrastructure Partners (Bolder investor)", "LinkedIn first", "IVP is Bolder Antwerp with an EU grant attached"),
    ("9", "Kreos - Aris Constantinides; Chicago Atlantic - John Mazarakis", "LinkedIn first; Greek-diaspora angle", "First European cannabis credit, secured on EU-GMP plant and Galenica receivables (2027)"),
    ("10", "Argus (Christos Akkelides), Muhanna (Marinos Theodosiou), Axia (Antonios Achilleoudis), J. Safra Sarasin (Rachel Koen)", "Warm", "ETYK / PEO IC decisions; Greek family-office process for IVP"),
    ("11", "Artemis Growth Partners; ELPEN; Cantourage / Curaleaf International", "LinkedIn / DP / commercial", "Strategic co-investment and offtake for Curity"),
    ("12", "Threadmark, Elm Capital, Cebile (one only)", "Cold", "IVP placement on success fee; no retainer above €25k without a named list"),
], [1.4, 6.4, 3.6, 6.4])

# ------------------------------------------------------------------ 9. Implications
doc.add_heading("9. Implications", level=1)
doc.add_heading("9.1 Regulatory", level=2)
bullet("Law 124(I)/2018 bars RAIFs from loan-originating strategies; AIFMD II (Directive (EU) 2024/927) adds leverage caps of 175% / 300% and 5% retention, and Cyprus has not fully transposed. Any note programme (7% Pyrum bond, 8% Curity note) must remain at Impactus Group Ltd or CPEF I Curity Ltd level, outside the RAIF's perimeter, and CySEC must be satisfied the AIFM's licence covers each strategy.", "CySEC / AIFMD II.")
bullet("Marketing to 'well-informed' investors is a Cyprus-law category; abroad they are non-professional and need the host national regime (German semi-professional €200k, Dutch €100k). Pre-marketing rules under Directive 2019/1160 apply to the IVP teaser.", "Marketing perimeter.")
bullet("HIIF's market-operator test and the Innovation Fund's cumulation rules govern whether the VD tranche is aid-free; a below-market return can itself fail the test. Development Law 4887/2022 aid intensity for Phthiotis is unverified.", "State aid.")
bullet("UK lenders and LPs require a POCA analysis on cannabis cash flows; US-parented platforms are constrained by federal illegality even for EU-GMP medical product. Segregate CPEF from any borrowing base; ring-fencing between compartments already bars cross-collateralisation.", "Cannabis.")
doc.add_heading("9.2 Tax", level=2)
bullet("Offtaker prepayments and deferred EPC payments are trade instruments, not loans, for Greek withholding and thin-capitalisation purposes; document the pricing under Cyprus Circular 6/2023 and OECD TPG (CUP) where Impactus Ltd invoices Thermo Lysi (the €401k outstanding invoice is a live transfer-pricing exposure).", "Structure.")
bullet("Warrants issued to a VD lender at ProjectCo level are equity instruments for Greek CIT (22%); interest on PIK accrues for deductibility subject to the 30% EBITDA interest-limitation rule; grant income is recognised as the milestones are accepted.", "Instruments.")
bullet("Cyprus 12.5% CIT and NID at IPL level remain unaffected; Pillar Two is not in scope at this group size.", "Cyprus.")
doc.add_heading("9.3 Liquidity", level=2)
bullet("The €5m IPL round funds development costs to FID; if FID slips to Q2 2027 the round covers about nine months of burn (ECI fees €0.86m, engineering, PMO, TDD). A second sponsor-level bridge should be planned now, not in March.", "Sponsor level.")
bullet("The PE Fund's 8% coupon and 2027 to 2028 operating costs (c.€3m per the August 2025 note) need a documented cash source before the provident-fund channel is scaled.", "Fund level.")
bullet("Investors in IPL rank behind the senior facility and any VD tranche at ProjectCo level; exit is a trade sale or platform exit within five years. State this in every approach.", "Investor position.")

# ------------------------------------------------------------------ 10. Data limits
doc.add_heading("10. What could not be verified", level=1)
bullet("The Cascade page and all firm websites, LinkedIn and press wires were blocked by the sandbox network policy; the session's web-search quota was exhausted after about 200 queries. Every external figure comes from search-engine extracts of the cited page, not a full read; re-verify before quoting externally.")
bullet("Apollo.io is on a free plan that blocks API search and enrichment; no contact data was purchased. External names are flagged 'verify' and carry LinkedIn search links, not profile URLs.")
bullet("Five logos on the map are illegible; candidates are flagged. Ask Cascade Debt for the underlying list.")
bullet("Curity figures rest on the FY2023 Meerkat plan; the FY2025 valuation and audited accounts should replace them before any lender sees the CPEF tab.")
bullet("The capex benchmark '100 vs 250' in the PMO minutes is recorded without units; the plant capex is unresolved between €118.5m and €150m until the firm EPC offer of 31 October 2026.")

p("")
p(DISCLAIMER, italic=True, size=7.5, color=GREY)

# footer page numbers
for section in doc.sections:
    f = section.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = f.add_run("Impactus Private Equity Group - Internal - Page "); run.font.size = Pt(8); run.font.color.rgb = GREY
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "1"; r.append(t); fld.append(r); f._p.append(fld)

doc.save(OUT); print("saved", OUT)
